from io import BytesIO
from pathlib import Path
from typing import Optional
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from fpdf import FPDF

from .document_utils import sanitize_text


BASE_DIR = Path(__file__).resolve().parents[1]
DEFAULT_LOGO = BASE_DIR / "assets" / "logo.png"


def _paragraphs(text: str):
    clean = sanitize_text(text)
    return [p.strip() for p in clean.split("\n") if p.strip()]


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def format_docx(text: str, doc_type: str, logo_path: Optional[str] = None) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    logo = Path(logo_path) if logo_path else DEFAULT_LOGO
    if logo.exists():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(logo), width=Inches(1.4))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(sanitize_text(doc_type).upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for paragraph_text in _paragraphs(text):
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(paragraph_text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        if re.match(r"^(\d+[.)]|[A-Z][A-Z ]{3,}|SECTION|ARTICLE|PARTIES|SIGNATURES)", paragraph_text):
            r.bold = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Generated with LegalEase - AI-Powered Legal Document Generator")
    fr.font.name = "Times New Roman"
    fr.font.size = Pt(8)

    out = BytesIO()
    document.save(out)
    return out.getvalue()


class LegalPDF(FPDF):
    def __init__(self, logo_path: Optional[str] = None):
        super().__init__()
        self.logo_path = str(logo_path) if logo_path else str(DEFAULT_LOGO)

    def header(self):
        if Path(self.logo_path).exists():
            self.image(self.logo_path, x=85, y=8, w=40)
            self.set_y(30)
        else:
            self.set_y(10)
        self.set_font("Helvetica", "B", 9)
        self.cell(0, 6, "LegalEase", align="C")
        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "", 8)
        self.cell(
            0,
            10,
            "Generated with LegalEase - Please review before legal use.",
            align="C",
        )


def format_pdf(text: str, doc_type: str, logo_path: Optional[str] = None) -> bytes:
    pdf = LegalPDF(logo_path)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 15)
    pdf.multi_cell(0, 8, sanitize_text(doc_type).upper(), align="C")
    pdf.ln(5)

    for paragraph_text in _paragraphs(text):
        is_heading = bool(
            re.match(
                r"^(\d+[.)]|[A-Z][A-Z ]{3,}|SECTION|ARTICLE|PARTIES|SIGNATURES)",
                paragraph_text,
            )
        )
        pdf.set_font("Helvetica", "B" if is_heading else "", 11)
        pdf.multi_cell(0, 6, paragraph_text)
        pdf.ln(2)

    return bytes(pdf.output())
